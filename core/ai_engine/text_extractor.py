import logging
import os
import re
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional
import fitz  # PyMuPDF

logger = logging.getLogger(__name__)

# Optional progress callback signature: (current_step: int, total_steps: int, message: str) -> None
ProgressCallback = Optional[Callable[[int, int, str], None]]

_ocr_engine = None


def _get_ocr_engine():
    """Lazy-loads RapidOCR only when needed."""
    global _ocr_engine
    if _ocr_engine is None:
        try:
            from rapidocr_onnxruntime import RapidOCR

            _ocr_engine = RapidOCR()
            logger.info("RapidOCR engine initialized for image-page extraction.")
        except Exception as exc:
            logger.warning(f"RapidOCR unavailable, falling back to digital text only: {exc}")
            _ocr_engine = False
    return _ocr_engine if _ocr_engine is not False else None


def _run_page_ocr(page: fitz.Page, dpi: int = 150) -> str:
    """
    Renders a PDF page at a CPU-friendly 150 DPI and extracts text via RapidOCR.
    """
    engine = _get_ocr_engine()
    if engine is None:
        return ""

    try:
        pix = page.get_pixmap(dpi=dpi, colorspace=fitz.csGRAY)
        img_bytes = pix.tobytes("png")
        result, _ = engine(img_bytes)
        if not result:
            return ""
        lines = [item[1] for item in result if len(item) >= 2 and isinstance(item[1], str)]
        return "\n".join(lines).strip()
    except Exception as exc:
        logger.debug(f"Page OCR failed on page {getattr(page, 'number', '?')}: {exc}")
        return ""


def extract_text_from_file(filepath: str) -> str:
    """
    High-performance native text extractor for documents.
    Supports: PDF, EPUB, DOCX, PPTX, TXT, MD.
    """
    if not filepath or not os.path.isfile(filepath):
        return ""

    path = Path(filepath)
    ext = path.suffix.lower()
    text = ""

    try:
        if ext in (".txt", ".md", ".csv", ".json", ".log"):
            with open(path, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()

        elif ext == ".pdf":
            extracted_pages = []
            try:
                with fitz.open(str(path)) as doc:
                    for i, page in enumerate(doc):
                        raw = page.get_text() or ""
                        clean = raw.replace("\x00", "").strip()
                        if len(clean) < 20:
                            # Pre-check: only run OCR if page actually contains embedded images (skips blank pages in 0ms)
                            if not page.get_images(full=False):
                                pass
                            else:
                                ocr_txt = _run_page_ocr(page)
                                if ocr_txt and len(ocr_txt) > len(clean):
                                    clean = ocr_txt
                        if clean:
                            extracted_pages.append(f"[Page {i+1}]\n{clean}")
            except Exception as fitz_err:
                logger.debug(f"PyMuPDF fallback to pypdf ({fitz_err}).")
                from pypdf import PdfReader
                reader = PdfReader(str(path))
                for i, page in enumerate(reader.pages):
                    extracted = page.extract_text()
                    if extracted and extracted.strip():
                        extracted_pages.append(f"[Page {i+1}]\n{extracted.strip()}")
            text = "\n\n".join(extracted_pages)

        elif ext == ".docx":
            import docx
            doc = docx.Document(str(path))
            parts = []
            for p in doc.paragraphs:
                if p.text.strip():
                    parts.append(p.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    row_txt = [c.text.strip() for c in row.cells if c.text.strip()]
                    if row_txt:
                        parts.append(" | ".join(row_txt))
            text = "\n\n".join(parts)

        elif ext in (".pptx", ".ppt"):
            from pptx import Presentation
            prs = Presentation(str(path))
            slides = []
            for i, slide in enumerate(prs.slides):
                slide_lines = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_lines.append(shape.text.strip())
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                    notes = slide.notes_slide.notes_text_frame.text.strip()
                    if notes:
                        slide_lines.append(f"[Notes: {notes}]")
                if slide_lines:
                    slides.append(f"--- Slide {i+1} ---\n" + "\n".join(slide_lines))
            text = "\n\n".join(slides)

        elif ext == ".epub":
            try:
                import ebooklib
                from ebooklib import epub
                from bs4 import BeautifulSoup
                book = epub.read_epub(str(path))
                parts = []
                for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT):
                    soup = BeautifulSoup(item.get_content(), "html.parser")
                    t = soup.get_text()
                    if t.strip():
                        parts.append(t.strip())
                text = "\n\n".join(parts)
            except Exception as epub_err:
                logger.warning(f"EPUB extraction fallback: {epub_err}")

    except Exception as exc:
        logger.error(f"Error extracting text from {path.name}: {exc}")
        return ""

    if text:
        text = text.replace("\x00", "").strip()
    return text


def extract_document_structure(
    file_path: str,
    enable_ocr: bool = True,
    max_ocr_pages: int = 15,
    ocr_min_avg_words: int = 25,
    progress_callback: ProgressCallback = None,
) -> List[Dict[str, Any]]:
    """
    Extracts structured page text from a PDF or text document with:
    1. Pre-check for embedded images before triggering CPU OCR (skips blank pages in 0ms).
    2. Adaptive OCR circuit breaker (stops OCR on 1,200-page comic/illustrated books).
    3. Live page-level telemetry callbacks so the UI never freezes.
    """
    if not file_path or not os.path.isfile(file_path):
        return []

    ext = os.path.splitext(file_path)[1].lower()
    if ext in (".txt", ".md", ".csv", ".json", ".log"):
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read().strip()
            return [
                {
                    "page": 1,
                    "page_number": 1,
                    "type": "document",
                    "text": content,
                    "source": "digital",
                }
            ] if content else []
        except Exception as exc:
            logger.error(f"Failed reading text file {file_path}: {exc}")
            return []

    if ext == ".epub":
        try:
            import ebooklib
            from ebooklib import epub
            from bs4 import BeautifulSoup
            book = epub.read_epub(file_path)
            sections = []
            idx = 1
            for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT):
                soup = BeautifulSoup(item.get_content(), "html.parser")
                t = soup.get_text()
                if t and t.strip():
                    clean = t.replace("\x00", "").strip()
                    sections.append({
                        "page": idx,
                        "page_number": idx,
                        "type": "chapter",
                        "text": clean,
                        "source": "digital",
                    })
                    idx += 1
            return sections
        except Exception as exc:
            logger.warning(f"Structured EPUB extraction notice: {exc}")

    if ext == ".docx":
        try:
            import docx
            doc = docx.Document(file_path)
            sections = []
            current_section = []
            sec_idx = 1
            for p in doc.paragraphs:
                if p.text.strip():
                    current_section.append(p.text.strip())
                if len(current_section) >= 20:
                    sections.append({
                        "page": sec_idx,
                        "page_number": sec_idx,
                        "type": "section",
                        "text": "\n".join(current_section),
                        "source": "digital",
                    })
                    sec_idx += 1
                    current_section = []
            if current_section:
                sections.append({
                    "page": sec_idx,
                    "page_number": sec_idx,
                    "type": "section",
                    "text": "\n".join(current_section),
                    "source": "digital",
                })
            return sections
        except Exception as exc:
            logger.warning(f"Structured DOCX extraction notice: {exc}")

    if ext in (".pptx", ".ppt"):
        try:
            from pptx import Presentation
            prs = Presentation(file_path)
            sections = []
            for i, slide in enumerate(prs.slides):
                slide_lines = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_lines.append(shape.text.strip())
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                    notes = slide.notes_slide.notes_text_frame.text.strip()
                    if notes:
                        slide_lines.append(f"[Notes: {notes}]")
                if slide_lines:
                    clean = "\n".join(slide_lines).replace("\x00", "").strip()
                    sections.append({
                        "page": i + 1,
                        "page_number": i + 1,
                        "type": "slide",
                        "text": clean,
                        "source": "digital",
                    })
            return sections
        except Exception as exc:
            logger.warning(f"Structured PPTX extraction notice: {exc}")

    pages_output: List[Dict[str, Any]] = []
    ocr_pages_used = 0
    ocr_words_extracted = 0
    ocr_circuit_tripped = False

    try:
        with fitz.open(file_path) as doc:
            total_pages = len(doc)
            if total_pages == 0:
                return []

            for page_idx in range(total_pages):
                page = doc[page_idx]
                page_num = page_idx + 1

                # Emit periodic telemetry every 10 digital pages so the UI progress bar stays alive
                if progress_callback and (page_num == 1 or page_num % 10 == 0 or page_num == total_pages):
                    progress_callback(
                        page_num,
                        total_pages,
                        f"Extracting Page {page_num}/{total_pages}...",
                    )

                raw_text = page.get_text("text") or ""
                clean_text = re.sub(r"\s+", " ", raw_text).replace("\x00", "").strip()
                source_type = "digital"

                # Only consider OCR if digital text is < 20 chars AND OCR is allowed
                if len(clean_text) < 20 and enable_ocr and not ocr_circuit_tripped:
                    # Check 1: Does this page actually have an embedded image?
                    # Blank spacer pages have 0 images -> skip immediately!
                    page_images = page.get_images(full=False)
                    if not page_images:
                        # Blank page - do not trigger CPU OCR
                        pass
                    else:
                        if ocr_pages_used >= max_ocr_pages:
                            ocr_circuit_tripped = True
                            logger.info(
                                f"OCR budget ({max_ocr_pages} pages) reached for '{os.path.basename(file_path)}'. "
                                "Switching remaining pages to fast digital extraction."
                            )
                        else:
                            if progress_callback:
                                progress_callback(
                                    page_num,
                                    total_pages,
                                    f"Running OCR on page {page_num}/{total_pages} ({ocr_pages_used + 1}/{max_ocr_pages} max)...",
                                )
                            ocr_txt = _run_page_ocr(page)
                            ocr_pages_used += 1
                            word_count = len(ocr_txt.split())
                            ocr_words_extracted += word_count

                            if len(ocr_txt) >= 20:
                                clean_text = ocr_txt
                                source_type = "ocr"

                            # Check 2: Low-density circuit breaker (e.g., comic books / graphic novels)
                            # After 5 OCR pages, if average yield is under 25 words per page, trip circuit breaker
                            if ocr_pages_used >= 5:
                                avg_words = ocr_words_extracted / ocr_pages_used
                                if avg_words < ocr_min_avg_words:
                                    ocr_circuit_tripped = True
                                    logger.info(
                                        f"Low OCR text density ({avg_words:.1f} words/page) detected in "
                                        f"'{os.path.basename(file_path)}'. Bypassing OCR for remaining pages."
                                    )

                if len(clean_text) >= 20:
                    pages_output.append(
                        {
                            "page": page_num,
                            "page_number": page_num,
                            "type": "page",
                            "text": clean_text,
                            "source": source_type,
                        }
                    )

    except Exception as exc:
        logger.error(f"PDF extraction error on '{file_path}': {exc}")
        try:
            from pypdf import PdfReader
            reader = PdfReader(file_path)
            total_pdf_pages = len(reader.pages)
            for i, page in enumerate(reader.pages):
                page_num = i + 1
                if progress_callback and (page_num == 1 or page_num % 10 == 0 or page_num == total_pdf_pages):
                    progress_callback(page_num, total_pdf_pages, f"Extracting Page {page_num}/{total_pdf_pages}...")
                extracted = page.extract_text()
                if extracted and extracted.strip():
                    clean = extracted.replace("\x00", "").strip()
                    if len(clean) >= 20:
                        pages_output.append({
                            "page": page_num,
                            "page_number": page_num,
                            "type": "page",
                            "text": clean,
                            "source": "digital",
                        })
        except Exception as pypdf_err:
            logger.debug(f"pypdf fallback error on '{file_path}': {pypdf_err}")

    return pages_output
