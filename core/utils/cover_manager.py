import os
import io
import re
import zipfile
import logging
import asyncio
from pathlib import Path
from typing import Optional, Dict, Any, List, Tuple
from PIL import Image, ImageDraw, ImageFont

logger = logging.getLogger("COVER_MANAGER")

class CoverManager:
    """
    Optimized, universal Cover Extraction and Generation Service for Libre-Library.
    
    Supports:
      - PDF: PyMuPDF (fitz) page 0 rendering to high-res, optimized JPEG.
      - EPUB: Embedded cover art extraction (ebooklib/zip) with PyMuPDF page 0 fallback.
      - PPTX: Embedded slide thumbnail extraction with presentation title card fallback.
      - DOCX: Embedded document thumbnail with typographic front cover fallback.
      - TXT / MD / Fallback: Publication-grade typographic book cover generation.
    """
    
    TARGET_WIDTH = 600
    TARGET_HEIGHT = 900
    JPEG_QUALITY = 85

    @classmethod
    def _optimize_and_save_image(cls, pil_img: Image.Image, output_path: Path) -> bool:
        """Standardizes, resizes, and saves any PIL image to an optimized progressive JPEG."""
        try:
            # Handle RGBA / transparency gracefully with a white/clean backdrop
            if pil_img.mode in ('RGBA', 'LA') or (pil_img.mode == 'P' and 'transparency' in pil_img.info):
                alpha = pil_img.convert('RGBA')
                bg = Image.new('RGB', alpha.size, (255, 255, 255))
                bg.paste(alpha, mask=alpha.split()[3])
                img = bg
            elif pil_img.mode != 'RGB':
                img = pil_img.convert('RGB')
            else:
                img = pil_img

            # Preserve aspect ratio while constraining to maximum cover dimensions
            img.thumbnail((cls.TARGET_WIDTH, cls.TARGET_HEIGHT), Image.Resampling.LANCZOS)

            output_path.parent.mkdir(parents=True, exist_ok=True)
            img.save(str(output_path), format='JPEG', quality=cls.JPEG_QUALITY, optimize=True, progressive=True)
            return True
        except Exception as e:
            logger.warning(f"Error optimizing image for {output_path.name}: {e}")
            return False

    @classmethod
    def _extract_pdf_cover(cls, pdf_path: Path, output_path: Path) -> bool:
        """Renders page 0 of a PDF document at 150 DPI and saves as optimized JPEG."""
        doc = None
        try:
            import fitz
            doc = fitz.open(str(pdf_path))
            if len(doc) == 0:
                return False
            page = doc[0]
            pix = page.get_pixmap(dpi=150)
            
            mode = "RGBA" if pix.alpha else "RGB"
            img = Image.frombytes(mode, [pix.width, pix.height], pix.samples)
            return cls._optimize_and_save_image(img, output_path)
        except Exception as e:
            logger.warning(f"PDF cover extraction notice for {pdf_path.name}: {e}")
            return False
        finally:
            if doc is not None:
                try:
                    doc.close()
                except Exception:
                    pass

    @classmethod
    def _extract_epub_cover(cls, epub_path: Path, output_path: Path) -> bool:
        """
        Extracts embedded cover image from EPUB.
        Falls back to PyMuPDF page 0 rendering if no explicit cover item is found.
        """
        # Strategy 1: Check embedded cover image via ebooklib
        try:
            import ebooklib
            from ebooklib import epub
            book = epub.read_epub(str(epub_path))
            for item in book.get_items():
                if item.get_type() == ebooklib.ITEM_COVER or ('cover' in item.get_name().lower() and 'image' in (item.media_type or '')):
                    content = item.get_content()
                    if content and len(content) > 1000:
                        img = Image.open(io.BytesIO(content))
                        if cls._optimize_and_save_image(img, output_path):
                            return True
        except Exception as e:
            logger.debug(f"Epublib cover check note for {epub_path.name}: {e}")

        # Strategy 2: PyMuPDF natively opens and renders EPUB page 0
        doc = None
        try:
            import fitz
            doc = fitz.open(str(epub_path))
            if len(doc) > 0:
                page = doc[0]
                pix = page.get_pixmap(dpi=140)
                mode = "RGBA" if pix.alpha else "RGB"
                img = Image.frombytes(mode, [pix.width, pix.height], pix.samples)
                if cls._optimize_and_save_image(img, output_path):
                    return True
        except Exception as e:
            logger.debug(f"Fitz EPUB render note for {epub_path.name}: {e}")
        finally:
            if doc is not None:
                try:
                    doc.close()
                except Exception:
                    pass

        return False

    @classmethod
    def _extract_pptx_cover(cls, pptx_path: Path, output_path: Path, title: str = "", author: str = "") -> bool:
        """Extracts PPTX slide 1 thumbnail or generates an executive presentation card."""
        # Strategy 1: PowerPoint built-in thumbnail
        try:
            with zipfile.ZipFile(str(pptx_path)) as z:
                nl = z.namelist()
                if 'docProps/thumbnail.jpeg' in nl:
                    raw_thumb = z.read('docProps/thumbnail.jpeg')
                    img = Image.open(io.BytesIO(raw_thumb))
                    if cls._optimize_and_save_image(img, output_path):
                        return True
                elif 'docProps/thumbnail.png' in nl:
                    raw_thumb = z.read('docProps/thumbnail.png')
                    img = Image.open(io.BytesIO(raw_thumb))
                    if cls._optimize_and_save_image(img, output_path):
                        return True
        except Exception as e:
            logger.debug(f"PPTX thumbnail check notice: {e}")

        # Strategy 2: Extract slide 1 title & author for presentation card
        slide_title = title
        slide_author = author
        try:
            from pptx import Presentation
            prs = Presentation(str(pptx_path))
            if len(prs.slides) > 0:
                first_slide = prs.slides[0]
                texts = [shape.text.strip() for shape in first_slide.shapes if shape.has_text_frame and shape.text.strip()]
                if texts:
                    if not slide_title or slide_title == pptx_path.stem:
                        slide_title = texts[0].replace('\x0b', ' ').replace('\n', ' ').strip()
                    if len(texts) > 1 and (not slide_author or slide_author in ("Unknown", "Unknown Author")):
                        author_cand = texts[1].replace('\x0b', ' ').strip()
                        author_cand = re.sub(r'^(prepared by|members?|presented by|by)[:\s]*', '', author_cand, flags=re.IGNORECASE)
                        first_line = [l.strip() for l in author_cand.split('\n') if l.strip()]
                        if first_line:
                            slide_author = first_line[0]
        except Exception as e:
            logger.debug(f"PPTX slide text extraction notice: {e}")

        return cls._generate_typographic_cover(output_path, title=slide_title or pptx_path.stem, author=slide_author or "Author", file_type="PPTX")

    @classmethod
    def _extract_docx_cover(cls, docx_path: Path, output_path: Path, title: str = "", author: str = "") -> bool:
        """Extracts DOCX document thumbnail or generates a clean document cover."""
        try:
            with zipfile.ZipFile(str(docx_path)) as z:
                nl = z.namelist()
                if 'docProps/thumbnail.jpeg' in nl:
                    img = Image.open(io.BytesIO(z.read('docProps/thumbnail.jpeg')))
                    if cls._optimize_and_save_image(img, output_path):
                        return True
        except Exception as e:
            logger.debug(f"DOCX thumbnail notice: {e}")

        return cls._generate_typographic_cover(output_path, title=title or docx_path.stem, author=author or "Author", file_type="DOCX")

    @classmethod
    def _generate_typographic_cover(
        cls,
        output_path: Path,
        title: str,
        author: str = "Unknown",
        file_type: str = "DOCUMENT"
    ) -> bool:
        """
        Synthesizes a modern, publication-grade typographic book cover with subtle
        gradients, elegant borders, auto-wrapped typography, and metadata badges.
        """
        try:
            w, h = cls.TARGET_WIDTH, cls.TARGET_HEIGHT
            img = Image.new('RGB', (w, h))
            draw = ImageDraw.Draw(img)

            # Modern gradient background (Deep Midnight Indigo to Dark Slate)
            top_color = (20, 24, 48)
            bot_color = (15, 23, 42)
            for y in range(h):
                factor = y / h
                r = int(top_color[0] + factor * (bot_color[0] - top_color[0]))
                g = int(top_color[1] + factor * (bot_color[1] - top_color[1]))
                b = int(top_color[2] + factor * (bot_color[2] - top_color[2]))
                draw.line([(0, y), (w, y)], fill=(r, g, b))

            # Architectural borders
            draw.rectangle([(22, 22), (w - 22, h - 22)], outline=(51, 65, 85), width=2)
            draw.rectangle([(30, 30), (w - 30, h - 30)], outline=(99, 102, 241), width=1)
            draw.rectangle([(38, 38), (w - 38, h - 38)], outline=(30, 41, 59), width=1)

            # Resolve Windows system fonts with safe fallbacks
            def _get_font(candidates: List[str], size: int) -> ImageFont.ImageFont:
                for font_name in candidates:
                    cand = Path(f"C:/Windows/Fonts/{font_name}")
                    if cand.exists():
                        try:
                            return ImageFont.truetype(str(cand), size)
                        except Exception:
                            pass
                try:
                    return ImageFont.load_default()
                except Exception:
                    return None

            badge_font = _get_font(["segoeuib.ttf", "arialbd.ttf"], 13)
            author_font = _get_font(["segoeui.ttf", "arial.ttf"], 22)
            footer_font = _get_font(["segoeuib.ttf", "arialbd.ttf"], 11)

            # Header Badge
            clean_fmt = file_type.upper().replace('.', '')
            badge_text = f"LIBRE ARCHIVE  •  {clean_fmt}"
            draw.text((w // 2, 85), badge_text, font=badge_font, fill=(165, 180, 252), anchor='mm')
            draw.line([(w // 2 - 60, 108), (w // 2 + 60, 108)], fill=(99, 102, 241), width=2)

            # Dynamic Font Scaling for Title
            clean_title = re.sub(r'[\r\n\t]+', ' ', title).strip() or "Untitled Document"
            title_size = 38 if len(clean_title) < 40 else (32 if len(clean_title) < 80 else 26)
            title_font = _get_font(["georgiab.ttf", "segoeuib.ttf", "arialbd.ttf"], title_size)

            # Smart line wrap
            words = clean_title.split()
            lines = []
            cur_line = ""
            max_line_width = w - 100
            for word in words:
                test_line = f"{cur_line} {word}".strip()
                bbox = draw.textbbox((0, 0), test_line, font=title_font)
                if (bbox[2] - bbox[0]) <= max_line_width:
                    cur_line = test_line
                else:
                    if cur_line:
                        lines.append(cur_line)
                    cur_line = word
            if cur_line:
                lines.append(cur_line)

            lines = lines[:6] # Max 6 lines
            line_height = int(title_size * 1.35)
            total_text_height = len(lines) * line_height
            
            # Position title vertically in upper-middle area
            y_start = max(180, 360 - (total_text_height // 2))
            for i, line in enumerate(lines):
                draw.text((w // 2, y_start + (i * line_height)), line, font=title_font, fill=(248, 250, 252), anchor='mm')

            # Accent separator rule
            sep_y = y_start + total_text_height + 25
            draw.line([(w // 2 - 45, sep_y), (w // 2 + 45, sep_y)], fill=(129, 140, 248), width=2)

            # Author attribution
            clean_author = str(author or "Unknown Author").strip()
            if clean_author.lower() in ("unknown", "unknown author", "none"):
                clean_author = "Unknown Author"
            draw.text((w // 2, sep_y + 40), f"by {clean_author}", font=author_font, fill=(203, 213, 225), anchor='mm')

            # Bottom decorative brand badge
            draw.line([(w // 2 - 90, h - 85), (w // 2 + 90, h - 85)], fill=(51, 65, 85), width=1)
            draw.text((w // 2, h - 65), "DIGITAL ARCHIVE EDITION", font=footer_font, fill=(100, 116, 139), anchor='mm')

            output_path.parent.mkdir(parents=True, exist_ok=True)
            img.save(str(output_path), format='JPEG', quality=cls.JPEG_QUALITY, optimize=True, progressive=True)
            return True
        except Exception as e:
            logger.warning(f"Error generating typographic cover: {e}")
            return False

    @classmethod
    def extract_or_generate_cover(
        cls,
        file_path: Path,
        output_path: Path,
        title: str = "",
        author: str = "",
        file_type: str = ""
    ) -> bool:
        """
        Synchronously extracts or generates a cover image for the given document file.
        Dispatches to the optimal strategy for the file format.
        """
        if not file_path.exists() or file_path.stat().st_size == 0:
            return False

        ext = file_path.suffix.lower()
        success = False

        if ext == '.pdf':
            success = cls._extract_pdf_cover(file_path, output_path)
        elif ext == '.epub':
            success = cls._extract_epub_cover(file_path, output_path)
        elif ext in ('.pptx', '.ppt'):
            success = cls._extract_pptx_cover(file_path, output_path, title=title, author=author)
        elif ext in ('.docx', '.doc'):
            success = cls._extract_docx_cover(file_path, output_path, title=title, author=author)
        elif ext in ('.txt', '.md'):
            success = cls._generate_typographic_cover(
                output_path,
                title=title or file_path.stem.replace('_', ' ').replace('-', ' ').title(),
                author=author or "Unknown Author",
                file_type="TXT" if ext == '.txt' else "MARKDOWN"
            )

        # Resilient Fallback: If format-specific extractor could not extract a cover,
        # generate an elegant typographic front cover so the document never displays a missing placeholder.
        if not success or not output_path.exists() or output_path.stat().st_size == 0:
            success = cls._generate_typographic_cover(
                output_path,
                title=title or file_path.stem.replace('_', ' ').replace('-', ' ').title(),
                author=author or "Unknown Author",
                file_type=file_type or ext.replace('.', '').upper() or "DOCUMENT"
            )

        return success

    @classmethod
    async def extract_cover_async(
        cls,
        file_path: Path,
        output_path: Path,
        title: str = "",
        author: str = "",
        file_type: str = ""
    ) -> bool:
        """Non-blocking asynchronous wrapper running cover extraction in the default thread pool."""
        return await asyncio.to_thread(
            cls.extract_or_generate_cover,
            file_path,
            output_path,
            title,
            author,
            file_type
        )

    @classmethod
    async def batch_generate_missing_covers(
        cls,
        books_dir: Optional[Path] = None,
        sync_mongodb: bool = True
    ) -> Dict[str, Any]:
        """
        Scans data/books, identifies all books missing a cover.jpg, and generates
        authentic front page covers or typographic cards. Synchronizes MongoDB.
        """
        from core.config import settings
        from core.database.mongo_manager import mongo_db
        import json

        base_dir = books_dir or (settings.BASE_DIR / 'data' / 'books')
        if not base_dir.exists():
            return {"total": 0, "success": 0, "failed": 0, "skipped": 0}

        results = {"total": 0, "success": 0, "failed": 0, "skipped": 0}
        candidate_folders: List[Path] = []

        for folder in sorted(base_dir.iterdir()):
            if not folder.is_dir():
                continue
            book_id = folder.name
            cover = folder / 'cover.jpg'
            
            # Find content file
            content_files = [
                f for f in folder.iterdir()
                if f.is_file() and f.suffix.lower() in ('.pdf', '.epub', '.pptx', '.ppt', '.docx', '.doc', '.txt', '.md')
            ]
            content_file = content_files[0] if content_files else None
            ext = content_file.suffix.lower() if content_file else ""
            file_type_str = ext.replace('.', '').upper() if ext else "BOOK"

            if cover.exists() and cover.stat().st_size > 500:
                results["skipped"] += 1
                # Sync MongoDB record if cover_image or file_type is missing
                if sync_mongodb and mongo_db.db is not None:
                    try:
                        sync_fields = {"cover_image": f"/static_books/{book_id}/cover.jpg"}
                        if content_file:
                            sync_fields["local_path"] = str(content_file)
                        if file_type_str != "BOOK":
                            sync_fields["file_type"] = file_type_str

                        await mongo_db.db['books'].update_one(
                            {"id": book_id, "$or": [{"cover_image": None}, {"cover_image": "/static/default_cover.svg"}, {"file_type": "UNKNOWN"}, {"file_type": None}]},
                            {"$set": sync_fields}
                        )
                    except Exception as sync_err:
                        logger.debug(f"Sync check notice for {book_id}: {sync_err}")
            else:
                candidate_folders.append(folder)

        results["total"] = len(candidate_folders)
        if not candidate_folders:
            return results

        logger.info(f"Starting batch cover generation for {len(candidate_folders)} books...")

        for folder in candidate_folders:
            book_id = folder.name
            cover_path = folder / 'cover.jpg'
            
            # Find content file
            content_files = [
                f for f in folder.iterdir()
                if f.is_file() and f.suffix.lower() in ('.pdf', '.epub', '.pptx', '.ppt', '.docx', '.doc', '.txt', '.md')
            ]
            if not content_files:
                results["failed"] += 1
                continue

            content_file = content_files[0]
            ext = content_file.suffix.lower()

            # Attempt to read title & author from metadata.json if present
            meta_path = folder / 'metadata.json'
            title = ""
            author = ""
            if meta_path.exists():
                try:
                    with open(meta_path, 'r', encoding='utf-8') as mf:
                        m_data = json.load(mf)
                        title = m_data.get('title', '')
                        authors_val = m_data.get('authors', [])
                        if isinstance(authors_val, list) and authors_val:
                            first_a = authors_val[0]
                            author = first_a.get('name', '') if isinstance(first_a, dict) else str(first_a)
                        elif isinstance(authors_val, str):
                            author = authors_val
                except Exception:
                    pass

            if not title:
                title = content_file.stem.replace('_', ' ').replace('-', ' ').title()

            file_type_str = ext.replace('.', '').upper()
            
            ok = await cls.extract_cover_async(
                content_file,
                cover_path,
                title=title,
                author=author,
                file_type=file_type_str
            )

            if ok and cover_path.exists():
                results["success"] += 1
                if sync_mongodb and mongo_db.db is not None:
                    try:
                        update_fields = {
                            "cover_image": f"/static_books/{book_id}/cover.jpg",
                            "local_path": str(content_file),
                            "file_type": file_type_str,
                        }
                        if title:
                            update_fields["title"] = title
                        if author and author != "Unknown":
                            update_fields["display_author"] = author

                        await mongo_db.db['books'].update_one(
                            {"id": book_id},
                            {"$set": update_fields},
                            upsert=False
                        )
                    except Exception as db_err:
                        logger.warning(f"MongoDB update notice for {book_id}: {db_err}")
            else:
                results["failed"] += 1

        logger.info(f"Batch cover generation complete: {results['success']} generated, {results['failed']} failed, {results['skipped']} previously existed.")
        return results

cover_manager = CoverManager()
