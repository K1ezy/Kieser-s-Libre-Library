import os
import re
import logging
import threading
from collections import OrderedDict
from pathlib import Path
from typing import Optional, List, Dict, Tuple, Any, Callable
import fitz  # PyMuPDF

logger = logging.getLogger("TEXT_EXTRACTOR")

# Optional progress callback signature: (current_step: int, total_steps: int, message: str) -> None
ProgressCallback = Optional[Callable[[int, int, str], None]]

# Thread-safe in-memory LRU cache for extracted text and document structure
_TEXT_CACHE: "OrderedDict[Tuple[str, float, int], str]" = OrderedDict()
_STRUCTURE_CACHE: "OrderedDict[Tuple[str, float, int, bool, int], List[dict]]" = OrderedDict()
_CACHE_LOCK = threading.Lock()
_MAX_CACHE_SIZE = 32

_OCR_ENGINE = None
_OCR_LOCK = threading.Lock()


def _get_ocr_engine():
    global _OCR_ENGINE
    if _OCR_ENGINE is None:
        with _OCR_LOCK:
            if _OCR_ENGINE is None:
                try:
                    from rapidocr_onnxruntime import RapidOCR
                    _OCR_ENGINE = RapidOCR()
                    logger.info("RapidOCR engine initialized for image-page extraction.")
                except Exception as e:
                    logger.debug(f"RapidOCR engine initialization note: {e}")
                    _OCR_ENGINE = False
    return _OCR_ENGINE if _OCR_ENGINE is not False else None


def _run_page_ocr(page: fitz.Page, dpi: int = 150) -> str:
    """Renders page pixmap at 150 DPI and runs RapidOCR on image-only or low-text PDF pages."""
    try:
        engine = _get_ocr_engine()
        if not engine:
            return ""
        pix = page.get_pixmap(dpi=dpi, colorspace=fitz.csGRAY)
        img_bytes = pix.tobytes("png")
        res, _ = engine(img_bytes)
        if res:
            lines = [line[1].strip() for line in res if line and len(line) > 1 and isinstance(line[1], str) and line[1].strip()]
            return "\n".join(lines).strip()
    except Exception as err:
        logger.debug(f"Page OCR extraction error: {err}")
    return ""


def extract_text_from_file(filepath: str) -> str:
    """
    High-performance native text extractor for documents.
    Supports: PDF, EPUB, DOCX, PPTX, TXT, MD.
    Thread-safe in-memory LRU caching accelerates repeated RAG/LLM passes.
    """
    path = Path(filepath)
    if not path.exists():
        return ""

    try:
        stat_res = path.stat()
        if stat_res.st_size == 0:
            return ""
        cache_key = (str(path.resolve()), stat_res.st_mtime, stat_res.st_size)
        with _CACHE_LOCK:
            if cache_key in _TEXT_CACHE:
                _TEXT_CACHE.move_to_end(cache_key)
                return _TEXT_CACHE[cache_key]
    except Exception:
        cache_key = None

    ext = path.suffix.lower()
    text = ""

    try:
        if ext == '.pdf':
            extracted_pages = []
            try:
                with fitz.open(str(path)) as doc:
                    for i, page in enumerate(doc):
                        t = page.get_text()
                        clean = t.replace('\x00', '').strip() if t else ""
                        if len(clean) < 20:
                            # Pre-check: Only run OCR if page actually contains embedded images (skips blank pages in 0ms)
                            if not page.get_images(full=False):
                                pass
                            else:
                                ocr_txt = _run_page_ocr(page)
                                if ocr_txt and len(ocr_txt) > len(clean):
                                    clean = ocr_txt
                        if clean:
                            extracted_pages.append(f"[Page {i+1}]\n{clean}")
            except Exception as fitz_err:
                logger.debug(f"PyMuPDF fallback to pypdf ({fitz_err}).")
                from pypdf import PdfReader
                reader = PdfReader(str(path))
                for i, page in enumerate(reader.pages):
                    extracted = page.extract_text()
                    if extracted and extracted.strip():
                        extracted_pages.append(f"[Page {i+1}]\n{extracted.strip()}")
            text = "\n\n".join(extracted_pages)

        elif ext == '.docx':
            import docx
            doc = docx.Document(str(path))
            parts = []
            for p in doc.paragraphs:
                if p.text.strip():
                    parts.append(p.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    row_txt = [c.text.strip() for c in row.cells if c.text.strip()]
                    if row_txt:
                        parts.append(" | ".join(row_txt))
            text = "\n\n".join(parts)

        elif ext in ('.pptx', '.ppt'):
            from pptx import Presentation
            prs = Presentation(str(path))
            slides = []
            for i, slide in enumerate(prs.slides):
                slide_lines = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_lines.append(shape.text.strip())
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                    notes = slide.notes_slide.notes_text_frame.text.strip()
                    if notes:
                        slide_lines.append(f"[Notes: {notes}]")
                if slide_lines:
                    slides.append(f"--- Slide {i+1} ---\n" + "\n".join(slide_lines))
            text = "\n\n".join(slides)

        elif ext == '.epub':
            try:
                import ebooklib
                from ebooklib import epub
                from bs4 import BeautifulSoup
                book = epub.read_epub(str(path))
                parts = []
                for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT):
                    soup = BeautifulSoup(item.get_content(), 'html.parser')
                    t = soup.get_text()
                    if t.strip():
                        parts.append(t.strip())
                text = "\n\n".join(parts)
            except Exception as epub_err:
                logger.warning(f"EPUB extraction fallback: {epub_err}")

        elif ext in ('.txt', '.md', '.csv', '.json', '.log'):
            with open(path, 'r', encoding='utf-8', errors='ignore') as f:
                text = f.read()

        else:
            logger.warning(f"Unsupported file extension for extraction: {ext}")
            return ""

    except Exception as e:
        logger.error(f"Error extracting text from {path.name}: {e}")
        return ""

    # Clean text: remove null bytes and strip excess whitespace
    if text:
        text = text.replace('\x00', '').strip()

    if cache_key and text:
        with _CACHE_LOCK:
            _TEXT_CACHE[cache_key] = text
            if len(_TEXT_CACHE) > _MAX_CACHE_SIZE:
                _TEXT_CACHE.popitem(last=False)

    return text


def extract_document_structure(
    filepath: str,
    enable_ocr: bool = True,
    max_ocr_pages: int = 15,
    ocr_min_avg_words: int = 25,
    progress_callback: ProgressCallback = None,
) -> List[Dict[str, Any]]:
    """
    Extracts structured document units (pages, slides, or chapters)
    with their respective page numbers.
    Features:
    1. Pre-check for embedded images before triggering CPU OCR (skips blank pages in 0ms).
    2. Adaptive OCR circuit breaker (stops OCR on 1,200-page comic/illustrated books).
    3. Live page-level telemetry callbacks so the UI never freezes.
    Returns: List of dicts [{"page": int, "page_number": int, "type": str, "text": str, "source": str}]
    """
    path = Path(filepath)
    if not path.exists():
        return []

    try:
        stat_res = path.stat()
        if stat_res.st_size == 0:
            return []
        cache_key = (str(path.resolve()), stat_res.st_mtime, stat_res.st_size, enable_ocr, max_ocr_pages)
        with _CACHE_LOCK:
            if cache_key in _STRUCTURE_CACHE:
                _STRUCTURE_CACHE.move_to_end(cache_key)
                return [dict(s) for s in _STRUCTURE_CACHE[cache_key]]
    except Exception:
        cache_key = None

    ext = path.suffix.lower()
    sections: List[Dict[str, Any]] = []

    try:
        if ext == '.pdf':
            ocr_pages_used = 0
            ocr_words_extracted = 0
            ocr_circuit_tripped = False
            try:
                with fitz.open(str(path)) as doc:
                    total_pages = len(doc)
                    for i, page in enumerate(doc):
                        page_num = i + 1

                        # Periodic telemetry every 10 digital pages so the UI stays alive
                        if progress_callback and (page_num == 1 or page_num % 10 == 0 or page_num == total_pages):
                            progress_callback(
                                page_num,
                                total_pages,
                                f"Extracting Page {page_num}/{total_pages}...",
                            )

                        t = page.get_text()
                        clean = re.sub(r"\s+", " ", t).replace('\x00', '').strip() if t else ""
                        source_type = "digital"

                        # Only consider OCR if digital text is < 20 chars AND OCR is allowed
                        if len(clean) < 20 and enable_ocr and not ocr_circuit_tripped:
                            # Check 1: Does this page actually have an embedded image?
                            # Blank spacer pages have 0 images -> skip immediately!
                            page_images = page.get_images(full=False)
                            if not page_images:
                                pass
                            else:
                                if ocr_pages_used >= max_ocr_pages:
                                    ocr_circuit_tripped = True
                                    logger.info(
                                        f"OCR budget ({max_ocr_pages} pages) reached for '{path.name}'. "
                                        "Switching remaining pages to fast digital extraction."
                                    )
                                else:
                                    if progress_callback:
                                        progress_callback(
                                            page_num,
                                            total_pages,
                                            f"Running OCR on page {page_num}/{total_pages} ({ocr_pages_used + 1}/{max_ocr_pages} max)...",
                                        )
                                    ocr_txt = _run_page_ocr(page)
                                    ocr_pages_used += 1
                                    words = len(ocr_txt.split())
                                    ocr_words_extracted += words

                                    if len(ocr_txt) >= 20:
                                        clean = ocr_txt
                                        source_type = "ocr"

                                    # Check 2: Low-density circuit breaker (after 5 OCR pages, average yield < 25 words/page)
                                    if ocr_pages_used >= 5:
                                        avg_words = ocr_words_extracted / ocr_pages_used
                                        if avg_words < ocr_min_avg_words:
                                            ocr_circuit_tripped = True
                                            logger.info(
                                                f"Low OCR text density ({avg_words:.1f} words/page) detected in "
                                                f"'{path.name}'. Bypassing OCR for remaining pages."
                                            )

                        if len(clean) >= 20:
                            sections.append({
                                "page": page_num,
                                "page_number": page_num,
                                "type": "page",
                                "text": clean,
                                "source": source_type,
                            })
            except Exception as fitz_err:
                logger.debug(f"PyMuPDF structured extraction fallback to pypdf: {fitz_err}")
                from pypdf import PdfReader
                reader = PdfReader(str(path))
                total_pdf_pages = len(reader.pages)
                for i, page in enumerate(reader.pages):
                    page_num = i + 1
                    if progress_callback and (page_num == 1 or page_num % 10 == 0 or page_num == total_pdf_pages):
                        progress_callback(page_num, total_pdf_pages, f"Extracting Page {page_num}/{total_pdf_pages}...")
                    extracted = page.extract_text()
                    if extracted and extracted.strip():
                        clean = extracted.replace('\x00', '').strip()
                        if len(clean) >= 20:
                            sections.append({
                                "page": page_num,
                                "page_number": page_num,
                                "type": "page",
                                "text": clean,
                                "source": "digital",
                            })

        elif ext in ('.pptx', '.ppt'):
            from pptx import Presentation
            prs = Presentation(str(path))
            for i, slide in enumerate(prs.slides):
                slide_lines = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_lines.append(shape.text.strip())
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                    notes = slide.notes_slide.notes_text_frame.text.strip()
                    if notes:
                        slide_lines.append(f"[Notes: {notes}]")
                if slide_lines:
                    clean = "\n".join(slide_lines).replace('\x00', '').strip()
                    sections.append({
                        "page": i + 1,
                        "page_number": i + 1,
                        "type": "slide",
                        "text": clean,
                        "source": "digital",
                    })

        elif ext == '.epub':
            try:
                import ebooklib
                from ebooklib import epub
                from bs4 import BeautifulSoup
                book = epub.read_epub(str(path))
                idx = 1
                for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT):
                    soup = BeautifulSoup(item.get_content(), 'html.parser')
                    t = soup.get_text()
                    if t and t.strip():
                        clean = t.replace('\x00', '').strip()
                        sections.append({
                            "page": idx,
                            "page_number": idx,
                            "type": "chapter",
                            "text": clean,
                            "source": "digital",
                        })
                        idx += 1
            except Exception as e:
                logger.warning(f"Structured EPUB extraction notice: {e}")

        elif ext == '.docx':
            import docx
            doc = docx.Document(str(path))
            current_section = []
            sec_idx = 1
            for p in doc.paragraphs:
                if p.text.strip():
                    current_section.append(p.text.strip())
                if len(current_section) >= 20:
                    sections.append({
                        "page": sec_idx,
                        "page_number": sec_idx,
                        "type": "section",
                        "text": "\n".join(current_section),
                        "source": "digital",
                    })
                    sec_idx += 1
                    current_section = []
            if current_section:
                sections.append({
                    "page": sec_idx,
                    "page_number": sec_idx,
                    "type": "section",
                    "text": "\n".join(current_section),
                    "source": "digital",
                })

        else:
            # Fallback for plain text, markdown, etc.
            raw = extract_text_from_file(filepath)
            if raw:
                lines = raw.split('\n')
                curr_lines = []
                curr_len = 0
                page_idx = 1
                for line in lines:
                    curr_lines.append(line)
                    curr_len += len(line)
                    if curr_len > 3000:
                        sections.append({
                            "page": page_idx,
                            "page_number": page_idx,
                            "type": "page",
                            "text": "\n".join(curr_lines),
                            "source": "digital",
                        })
                        page_idx += 1
                        curr_lines = []
                        curr_len = 0
                if curr_lines:
                    sections.append({
                        "page": page_idx,
                        "page_number": page_idx,
                        "type": "page",
                        "text": "\n".join(curr_lines),
                        "source": "digital",
                    })

    except Exception as e:
        logger.error(f"Error in extract_document_structure: {e}")

    # Fallback to single page if nothing segmented
    if not sections:
        full = extract_text_from_file(filepath)
        if full:
            sections.append({
                "page": 1,
                "page_number": 1,
                "type": "document",
                "text": full,
                "source": "digital",
            })

    if cache_key and sections:
        with _CACHE_LOCK:
            _STRUCTURE_CACHE[cache_key] = sections
            if len(_STRUCTURE_CACHE) > _MAX_CACHE_SIZE:
                _STRUCTURE_CACHE.popitem(last=False)

    return sections