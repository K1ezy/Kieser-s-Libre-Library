import asyncio
from nicegui import ui, app
from pathlib import Path
import urllib.parse
import re
import logging
from typing import Optional

from core.config import settings
from core.utils.security_utils import is_valid_book_id

logger = logging.getLogger("READER_INTERFACE")

class BookReader:
    def __init__(
        self,
        book_id: str,
        current_page: int = 1,
        total_pages: int = 1,
        user_id: str = "guest",
        user_role: str = "student",
        saved_cfi: Optional[str] = None,
        theme: str = "light",
        font_size: int = 16,
        font_family: str = "sans",
        line_height: float = 1.65
    ):
        self.book_id = str(book_id).strip()
        self.current_page = max(1, current_page)
        self.total_pages = max(1, total_pages)
        self.user_id = user_id
        self.user_role = str(user_role or "student").lower()
        self.saved_cfi = saved_cfi
        self.theme = theme if theme in ('light', 'sepia', 'dark', 'black') else 'light'
        self.font_size = max(12, min(32, int(font_size)))
        self.font_family = font_family
        self.line_height = line_height
        self.books_root = (settings.BASE_DIR / 'data' / 'books').resolve()
        self.book_path = None
        self.file_url = None
        self.file_type = None
        self.file_name = None
        self.error_msg = None
        self.pdf_iframe = None
        self.hud_label = None
        self.scroll_area = None
        self.content_container = None
        self.doc_card = None
        self.save_lock = asyncio.Lock()
        self._initial_saved = False

    def find_content_file(self) -> bool:
        """Security: Validates identifier and searches for supported format in book folder."""
        if not is_valid_book_id(self.book_id):
            self.error_msg = "Invalid book identifier format."
            return False

        candidate_path = (self.books_root / self.book_id).resolve()
        if candidate_path == self.books_root or not candidate_path.is_relative_to(self.books_root):
            self.error_msg = "Unauthorized document directory access."
            return False

        self.book_path = candidate_path
        if not self.book_path.exists():
            self.error_msg = f"Book folder not found on disk: {self.book_id}"
            return False

        files = sorted([f for f in self.book_path.iterdir() if f.is_file()])

        # 1. PDF
        for f in files:
            if f.suffix.lower() == '.pdf':
                self.file_name = f.name
                safe_name = urllib.parse.quote(f.name)
                self.file_url = f'/static_books/{self.book_id}/{safe_name}'
                self.file_type = 'pdf'
                return True

        # 2. EPUB
        for f in files:
            if f.suffix.lower() == '.epub':
                self.file_name = f.name
                safe_name = urllib.parse.quote(f.name)
                self.file_url = f'/static_books/{self.book_id}/{safe_name}'
                self.file_type = 'epub'
                return True

        # 3. DOCX
        for f in files:
            if f.suffix.lower() == '.docx':
                self.file_name = f.name
                safe_name = urllib.parse.quote(f.name)
                self.file_url = f'/static_books/{self.book_id}/{safe_name}'
                self.file_type = 'docx'
                return True

        # 4. PPTX
        for f in files:
            if f.suffix.lower() in ('.pptx', '.ppt'):
                self.file_name = f.name
                safe_name = urllib.parse.quote(f.name)
                self.file_url = f'/static_books/{self.book_id}/{safe_name}'
                self.file_type = 'pptx'
                return True

        # 5. TXT / MD
        for f in files:
            if f.suffix.lower() in ['.txt', '.md']:
                self.file_name = f.name
                safe_name = urllib.parse.quote(f.name)
                self.file_url = f'/static_books/{self.book_id}/{safe_name}'
                self.file_type = 'text'
                return True

        self.error_msg = "No readable file format found in book folder."
        return False

    async def save_progress(self, new_page: int, cfi: Optional[str] = None):
        """Asynchronously persists reading coordinate without interrupting the user."""
        new_page = max(1, min(self.total_pages, new_page))
        if new_page == self.current_page and self._initial_saved and not cfi:
            return
        self._initial_saved = True
        self.current_page = new_page
        if cfi:
            self.saved_cfi = cfi

        pct_val = round((new_page / max(self.total_pages, 1)) * 100.0, 1)
        status = 'completed' if (new_page >= self.total_pages and self.total_pages > 1) else 'reading'

        hud_text = f"Page {new_page} of {self.total_pages} ({pct_val:.0f}%)"
        if self.hud_label:
            self.hud_label.set_text(hud_text)

        # Live HUD update via DOM
        ui.run_javascript(f'''
            const hud = document.getElementById("reader-hud-text");
            if (hud) {{
                hud.innerText = "{hud_text}";
            }}
        ''')

        async with self.save_lock:
            try:
                from core.database.mongo_manager import mongo_db
                await mongo_db.update_book_reading_status(self.book_id, status)
                if self.user_id:
                    await mongo_db.progress.save_reading_progress(
                        user_id=self.user_id,
                        book_id=self.book_id,
                        current_page=new_page,
                        total_pages=self.total_pages,
                        status=status,
                        cfi=self.saved_cfi
                    )
            except Exception as e:
                logger.error(f"Save progress error: {e}")

    async def step_page(self, delta: int):
        """Advances or rewinds reading position by delta and syncs display."""
        action = 'next' if delta > 0 else 'prev'
        if self.file_type == 'epub':
            # Authoritative reader navigation: send command into EPUB viewer
            ui.run_javascript(f'''
                const frame = document.getElementById("reader-epub-frame");
                if (frame && frame.contentWindow) {{
                    frame.contentWindow.postMessage({{ type: "READER_NAVIGATE", action: "{action}" }}, "*");
                }}
            ''')
            return
        elif self.file_type == 'pdf':
            # Authoritative reader navigation: send command into PDF.js viewer
            ui.run_javascript(f'''
                const frame = document.getElementById("reader-pdf-frame");
                if (frame && frame.contentWindow) {{
                    frame.contentWindow.postMessage({{ type: "READER_NAVIGATE", action: "{action}" }}, "*");
                }}
            ''')
            return

        new_p = max(1, min(self.total_pages, self.current_page + delta))
        if new_p != self.current_page:
            await self.save_progress(new_p)
            if self.file_type in ('docx', 'pptx', 'text') and self.scroll_area and self.total_pages > 1:
                target_pct = (new_p - 1) / (self.total_pages - 1)
                self.scroll_area.scroll_to(percent=target_pct)

    async def handle_scroll(self, e):
        """Calculates page position using normalized scroll ratio: scroll_y / max_scroll_y."""
        pct = getattr(e, 'vertical_percentage', None)
        if pct is not None:
            calculated_p = max(1, min(self.total_pages, round(pct * (self.total_pages - 1)) + 1))
            if calculated_p != self.current_page:
                await self.save_progress(calculated_p)

    def get_prose_theme_classes(self):
        """Returns container styling and background colors based on active theme."""
        if self.theme == 'sepia':
            return 'bg-[#fbf0d9] text-[#3d2f1d] border-[#d9c8af]', '#f4ecd8'
        elif self.theme == 'dark':
            return 'bg-[#1e293b] text-[#e2e8f0] border-slate-700', '#0f172a'
        elif self.theme == 'black':
            return 'bg-[#09090b] text-[#d1d5db] border-zinc-800', '#000000'
        return 'bg-white text-slate-800 border-slate-200', '#f8fafc'

    async def update_prose_theme(self, new_theme: str):
        """Updates theme for native document formats and persists preference."""
        self.theme = new_theme
        if self.user_id:
            try:
                from core.database.mongo_manager import mongo_db
                await mongo_db.save_reader_preferences(self.user_id, {"theme": new_theme})
                await mongo_db.progress.save_reading_progress(
                    user_id=self.user_id,
                    book_id=self.book_id,
                    current_page=self.current_page,
                    total_pages=self.total_pages,
                    status='reading',
                    theme=new_theme
                )
            except Exception:
                pass
        
        card_cls, bg_color = self.get_prose_theme_classes()
        ui.run_javascript(f'''
            const container = document.getElementById("reader-content-container");
            if (container) container.style.backgroundColor = "{bg_color}";
            const card = document.getElementById("prose-card-container");
            if (card) {{
                card.className = "w-full h-full p-4 sm:p-8 md:p-16 max-w-4xl shadow-sm sm:rounded-2xl sm:my-4 border transition-colors {card_cls}";
            }}
        ''')

    async def update_prose_font_size(self, delta: int):
        """Adjusts font size for native document formats."""
        self.font_size = max(12, min(32, self.font_size + delta))
        if self.user_id:
            try:
                from core.database.mongo_manager import mongo_db
                await mongo_db.save_reader_preferences(self.user_id, {"font_size": self.font_size})
            except Exception:
                pass
        ui.run_javascript(f'''
            const card = document.getElementById("prose-card-container");
            if (card) {{
                card.style.fontSize = "{self.font_size}px";
            }}
            const fontLabel = document.getElementById("prose-font-size-label");
            if (fontLabel) fontLabel.innerText = "{self.font_size}px";
        ''')

    async def build_ui(self):
        found = self.find_content_file()
        card_cls, container_bg = self.get_prose_theme_classes()

        # HTML Head Bridge: Live HUD synchronization, theme bridging, & keyboard routing
        ui.add_head_html('''
        <script>
        (function() {
            // Listen for authoritative reader progress & theme events from iframe (EPUB or PDF)
            window.addEventListener('message', function(event) {
                if (!event.data) return;
                if (event.data.type === 'LIBRE_READER_PROGRESS' || event.data.type === 'READER_PROGRESS_UPDATE') {
                    const page = event.data.current_page;
                    const total = event.data.total_pages;
                    const pct = event.data.percentage;
                    const hud = document.getElementById('reader-hud-text');
                    if (hud && page && total) {
                        hud.innerText = `Page ${page} of ${total} (${pct}%)`;
                    }
                }
                if (event.data.type === 'LIBRE_READER_THEME_CHANGE') {
                    const theme = event.data.theme;
                    const container = document.getElementById('reader-content-container');
                    if (container) {
                        if (theme === 'sepia') {
                            container.style.backgroundColor = '#f4ecd8';
                        } else if (theme === 'dark') {
                            container.style.backgroundColor = '#0f172a';
                        } else if (theme === 'black') {
                            container.style.backgroundColor = '#000000';
                        } else {
                            container.style.backgroundColor = '#f8fafc';
                        }
                    }
                }
            });

            // Keyboard navigation forwarder on parent window (routes to active EPUB or PDF reader)
            document.addEventListener('keydown', function(event) {
                if (['INPUT', 'TEXTAREA'].includes(event.target.tagName)) return;
                const epubFrame = document.getElementById('reader-epub-frame');
                const pdfFrame = document.getElementById('reader-pdf-frame');
                const frame = epubFrame || pdfFrame;
                if (event.key === 'ArrowRight' || event.key === 'PageDown') {
                    if (frame && frame.contentWindow) {
                        frame.contentWindow.postMessage({ type: 'READER_NAVIGATE', action: 'next' }, '*');
                    }
                } else if (event.key === 'ArrowLeft' || event.key === 'PageUp') {
                    if (frame && frame.contentWindow) {
                        frame.contentWindow.postMessage({ type: 'READER_NAVIGATE', action: 'prev' }, '*');
                    }
                }
            });
        })();
        </script>
        ''')

        # Fluid Non-Overlapping Header with Live HUD
        with ui.header().classes('bg-white/95 backdrop-blur-md border-b border-slate-200 text-slate-800 h-14 pt-[env(safe-area-inset-top,0px)] items-center px-2.5 sm:px-4 z-50 shadow-sm'):
            with ui.row().classes('w-full items-center justify-between flex-nowrap gap-1.5 sm:gap-2'):
                # Left: Back button & Title
                with ui.row().classes('items-center gap-1.5 flex-1 min-w-0'):
                    ui.button(icon='arrow_back', on_click=lambda: ui.navigate.to('/books')) \
                        .props('flat round dense color=indigo size=md').tooltip('Back to Library').classes('shrink-0')
                    ui.label(self.file_name or 'Reading Mode').classes('font-bold text-xs sm:text-sm text-slate-900 truncate flex-1 min-w-0')

                # Center: Live Reading Position HUD & Stepper Controls
                pct_val = round((self.current_page / max(self.total_pages, 1)) * 100.0, 1)
                with ui.row().classes('items-center gap-0.5 sm:gap-1.5 bg-slate-50 border border-slate-200 px-1.5 sm:px-3 py-1 rounded-xl shadow-2xs shrink-0'):
                    ui.button(icon='chevron_left', on_click=lambda: self.step_page(-1))\
                        .props('flat round dense size=xs color=indigo').tooltip('Previous Page')
                    
                    self.hud_label = ui.label(
                        f"Page {self.current_page} of {self.total_pages} ({pct_val:.0f}%)"
                    ).props('id=reader-hud-text').classes(
                        'text-[11px] sm:text-xs font-bold text-slate-700 min-w-0 sm:min-w-[130px] text-center font-mono inline-block'
                    )

                    ui.button(icon='chevron_right', on_click=lambda: self.step_page(1))\
                        .props('flat round dense size=xs color=indigo').tooltip('Next Page')

                # Right: Actions grouped and fixed (Enforces SRS NFR 18 Copyright Read-Only Protection)
                with ui.row().classes('items-center gap-1 sm:gap-1.5 shrink-0'):
                    # Appearance Menu for native documents (DOCX, PPTX, TXT)
                    if self.file_type in ('docx', 'pptx', 'text'):
                        with ui.button(icon='palette').props('flat round dense color=indigo size=sm sm:size=md').tooltip('Appearance & Themes'):
                            with ui.menu().classes('p-3 rounded-2xl shadow-xl border border-slate-200 bg-white min-w-[220px]'):
                                ui.label('READING THEME').classes('text-[10px] font-bold text-slate-400 uppercase tracking-wider mb-2')
                                with ui.row().classes('gap-1.5 mb-3'):
                                    ui.button('Light', on_click=lambda: self.update_prose_theme('light')).props('dense size=xs color=grey-3 text-color=dark')
                                    ui.button('Sepia', on_click=lambda: self.update_prose_theme('sepia')).props('dense size=xs color=amber-2 text-color=brown-9')
                                    ui.button('Dark', on_click=lambda: self.update_prose_theme('dark')).props('dense size=xs color=blue-grey-9 text-color=white')
                                    ui.button('Black', on_click=lambda: self.update_prose_theme('black')).props('dense size=xs color=dark text-color=white')
                                
                                ui.separator().classes('my-2')
                                ui.label('TEXT SIZE').classes('text-[10px] font-bold text-slate-400 uppercase tracking-wider mb-1')
                                with ui.row().classes('items-center justify-between'):
                                    ui.button('A-', on_click=lambda: self.update_prose_font_size(-2)).props('dense outline size=xs color=indigo')
                                    ui.label(f'{self.font_size}px').props('id=prose-font-size-label').classes('font-bold text-xs text-slate-700 font-mono')
                                    ui.button('A+', on_click=lambda: self.update_prose_font_size(2)).props('dense outline size=xs color=indigo')

                    if self.user_role in ('admin', 'librarian') and self.file_url:
                        ui.button(icon='download', on_click=lambda: ui.download(self.file_url)) \
                            .props('flat round dense color=indigo size=sm sm:size=md').tooltip('Librarian Export / Download')
                    else:
                        with ui.element('div').classes('flex items-center gap-1 px-1.5 sm:px-2.5 py-1 rounded-xl bg-slate-100 text-slate-500 text-[10px] sm:text-[11px] font-bold border border-slate-200').tooltip('Read-Only Copyright Protection (SRS NFR 18)'):
                            ui.icon('lock', size='13px').classes('text-slate-400')
                            ui.label('Read-Only').classes('hidden xs:inline')
                    ui.button(icon='fullscreen', on_click=lambda: ui.run_javascript('if(document.fullscreenElement){document.exitFullscreen()}else{document.documentElement.requestFullscreen()}')) \
                        .props('flat round dense color=grey-7 size=sm sm:size=md').tooltip('Toggle Fullscreen')

        with ui.column().props('id=reader-content-container').classes('w-full h-[calc(100dvh-3.5rem)] pt-[env(safe-area-inset-top,0px)] pb-[env(safe-area-inset-bottom,0px)] p-0 m-0 items-center justify-center overflow-hidden transition-colors').style(f'background-color: {container_bg}'):
            if not found:
                with ui.card().classes('items-center text-center p-6 sm:p-8 bg-white shadow-xl rounded-2xl sm:rounded-3xl border-l-4 border-rose-500 w-[calc(100vw-2rem)] max-w-md mx-4'):
                    ui.icon('error_outline', size='3.5em').classes('text-rose-500 mb-2')
                    ui.label("Document Unavailable").classes('text-base sm:text-lg font-bold text-slate-900')
                    ui.label(self.error_msg).classes('text-xs text-rose-500 font-mono mt-1 break-words')
                    ui.button('Back to Collection', on_click=lambda: ui.navigate.to('/books')).props('unelevated rounded color=indigo size=sm font-bold').classes('mt-4')

            elif self.file_type == 'pdf':
                viewer_path = (
                    f"/static/pdf-viewer/index.html?file={self.file_url}"
                    f"&book_id={self.book_id}&user_id={self.user_id}"
                    f"&page={self.current_page}&total={self.total_pages}"
                )
                self.pdf_iframe = ui.element('iframe').props(
                    f'id="reader-pdf-frame" src="{viewer_path}" frameborder="0"'
                ).classes('w-full h-full border-none bg-transparent')

            elif self.file_type == 'epub':
                quoted_cfi = urllib.parse.quote(self.saved_cfi) if self.saved_cfi else ""
                viewer_path = (
                    f"/static/epub-viewer/index.html?file={self.file_url}"
                    f"&book_id={self.book_id}&user_id={self.user_id}"
                    f"&cfi={quoted_cfi}&page={self.current_page}&total={self.total_pages}"
                )
                ui.element('iframe').props(
                    f'id="reader-epub-frame" src="{viewer_path}" frameborder="0"'
                ).classes('w-full h-full border-none bg-transparent')

            elif self.file_type == 'docx':
                try:
                    import docx
                    file_path = self.book_path / self.file_name
                    doc = await asyncio.to_thread(docx.Document, file_path)
                    self.scroll_area = ui.scroll_area().props('id=prose-card-container').classes(f'w-full h-full p-4 sm:p-8 md:p-16 max-w-4xl shadow-sm sm:rounded-2xl sm:my-4 border transition-colors {card_cls}').style(f'font-size: {self.font_size}px')
                    self.scroll_area.on_scroll(self.handle_scroll)
                    with self.scroll_area:
                        ui.label(f"Document: {self.file_name}").classes('text-xl sm:text-2xl font-bold mb-4 sm:mb-6 break-words')
                        for para in doc.paragraphs:
                            if para.text.strip():
                                ui.label(para.text).classes('font-serif mb-3 leading-relaxed break-words opacity-95')

                    if self.current_page > 1 and self.total_pages > 1:
                        target_pct = (self.current_page - 1) / (self.total_pages - 1)
                        ui.timer(0.4, lambda: self.scroll_area.scroll_to(percent=target_pct), once=True)
                except Exception as e:
                    ui.notify(f"Error reading DOCX: {e}", type='negative')

            elif self.file_type == 'pptx':
                try:
                    from pptx import Presentation
                    file_path = self.book_path / self.file_name
                    prs = await asyncio.to_thread(Presentation, file_path)
                    self.scroll_area = ui.scroll_area().props('id=prose-card-container').classes(f'w-full h-full p-3 sm:p-8 md:p-16 max-w-5xl shadow-sm sm:rounded-2xl sm:my-4 border transition-colors {card_cls}')
                    self.scroll_area.on_scroll(self.handle_scroll)
                    with self.scroll_area:
                        for i, slide in enumerate(prs.slides):
                            with ui.card().classes('w-full mb-4 sm:mb-6 p-4 sm:p-6 rounded-xl sm:rounded-2xl shadow-sm border border-inherit bg-inherit'):
                                ui.label(f"Slide {i+1}").classes('text-xs font-bold text-indigo-500 uppercase mb-2')
                                txts = [s.text for s in slide.shapes if hasattr(s, "text") and s.text.strip()]
                                if txts:
                                    ui.label(txts[0]).classes('text-base sm:text-lg font-bold mb-2 sm:mb-3 break-words')
                                    for t in txts[1:]: ui.markdown(f"• {t}").classes('ml-2 sm:ml-4 text-xs sm:text-sm break-words opacity-90')

                    if self.current_page > 1 and self.total_pages > 1:
                        target_pct = (self.current_page - 1) / (self.total_pages - 1)
                        ui.timer(0.4, lambda: self.scroll_area.scroll_to(percent=target_pct), once=True)
                except Exception as e:
                    ui.notify(f"Error reading PPTX: {e}", type='negative')

            elif self.file_type == 'text':
                try:
                    file_path = self.book_path / self.file_name
                    def _read_txt(fp):
                        with open(fp, 'r', encoding='utf-8', errors='ignore') as f:
                            return f.read()
                    content_txt = await asyncio.to_thread(_read_txt, file_path)
                    self.scroll_area = ui.scroll_area().props('id=prose-card-container').classes(f'w-full h-full p-4 sm:p-8 border transition-colors {card_cls}').style(f'font-size: {self.font_size}px')
                    self.scroll_area.on_scroll(self.handle_scroll)
                    with self.scroll_area:
                        ui.markdown(content_txt).classes('prose max-w-none lg:prose-xl mx-auto font-serif p-4 sm:p-8 break-words text-inherit')

                    if self.current_page > 1 and self.total_pages > 1:
                        target_pct = (self.current_page - 1) / (self.total_pages - 1)
                        ui.timer(0.4, lambda: self.scroll_area.scroll_to(percent=target_pct), once=True)
                except Exception as e:
                    ui.notify(f"Error reading text: {e}", type='negative')


async def reader_page(book_id: str):
    app.storage.client['page_path'] = '/read'
    from core.database.mongo_manager import mongo_db

    user_id = app.storage.user.get('user_id') or app.storage.user.get('token') or 'guest'
    from core.auth.jwt_handler import get_current_user_role
    user_role = get_current_user_role()
    book_doc = await mongo_db.get_book_details(book_id)
    total_pages = await mongo_db.get_or_resolve_page_count(book_id, book_doc)
    total_pages = max(1, total_pages)

    cur_p = 1
    saved_cfi = None
    existing_prog = None
    pref_theme = "light"
    pref_font_size = 16
    pref_font_family = "sans"
    pref_line_height = 1.65

    if user_id:
        user_prefs = await mongo_db.get_reader_preferences(user_id)
        if user_prefs:
            pref_theme = user_prefs.get('theme', pref_theme)
            pref_font_size = user_prefs.get('font_size', pref_font_size)
            pref_font_family = user_prefs.get('font_family', pref_font_family)
            pref_line_height = user_prefs.get('line_height', pref_line_height)

        existing_prog = await mongo_db.progress.get_reading_progress(user_id, book_id)
        if existing_prog:
            if existing_prog.get('current_page'):
                cur_p = int(existing_prog['current_page'])
            if existing_prog.get('cfi'):
                saved_cfi = str(existing_prog['cfi'])
            if existing_prog.get('theme'):
                pref_theme = existing_prog['theme']
            if existing_prog.get('font_size'):
                pref_font_size = existing_prog['font_size']
            if existing_prog.get('font_family'):
                pref_font_family = existing_prog['font_family']
            if existing_prog.get('line_height'):
                pref_line_height = existing_prog['line_height']

    cur_p = max(1, min(cur_p, total_pages))

    reader = BookReader(
        book_id=book_id,
        current_page=cur_p,
        total_pages=total_pages,
        user_id=user_id,
        user_role=user_role,
        saved_cfi=saved_cfi,
        theme=pref_theme,
        font_size=pref_font_size,
        font_family=pref_font_family,
        line_height=pref_line_height
    )
    await reader.build_ui()

    # Only initialize a brand-new reading entry if the user has no prior history with this book
    if not existing_prog:
        await reader.save_progress(cur_p)